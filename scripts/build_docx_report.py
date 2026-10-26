"""Build an editable Word version of the current technical report."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "technical-report.md"
OUTPUT = ROOT / "output" / "docx" / "technical-report.docx"
FONT = "Microsoft YaHei"


def set_run_font(run, size=None, bold=None, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_rich_text(paragraph, text, base_size=9.4):
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for token in tokens:
        if not token:
            continue
        code = token.startswith("`") and token.endswith("`")
        bold = token.startswith("**") and token.endswith("**")
        value = token[1:-1] if code else token[2:-2] if bold else token
        run = paragraph.add_run(value)
        set_run_font(run, size=base_size, bold=bold)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8, color="555555")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(9.4)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for style_name, size, before, after in (("Heading 1", 13, 10, 5), ("Heading 2", 10, 7, 3)):
    style = doc.styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title_style = doc.styles["Title"]
title_style.font.name = FONT
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(4)
title_style.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_before = Pt(0)
add_page_field(footer)

for block in SOURCE.read_text(encoding="utf-8").split("\n\n"):
    block = block.strip()
    if not block:
        continue
    if block.startswith("# "):
        p = doc.add_paragraph(style="Title")
        add_rich_text(p, block[2:], 22)
    elif block.startswith("## "):
        p = doc.add_paragraph(style="Heading 1")
        if block.startswith("## 实测结果"):
            p.paragraph_format.page_break_before = True
        add_rich_text(p, block[3:], 13)
    elif block.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.22)
        p.paragraph_format.space_after = Pt(3)
        add_rich_text(p, block[2:])
    else:
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        add_rich_text(p, block.replace("\n", " "))

doc.core_properties.title = "化工反应器建模 Agent 技术报告"
doc.core_properties.subject = "技术选型、系统架构、HYSYS 自动化探索、问题处理与验证边界"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
